"""Generate the blank form templates the PDF export fills in.

    python tools/make_templates.py

Writes:
    forms/swp/swp-form.pdf     fillable Safe Work Permit (AcroForm)
    forms/loto/loto-form.xlsx  LOTO workbook (Part I boundaries, Part II verifiers)

Field names in the SWP PDF match the swp_permits column names, so app/pdf.py
can map database rows straight onto the form. Swap these files for your own
site's forms by keeping the same field names / cell addresses.

Needs reportlab (dev only): pip install reportlab
"""
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.worksheet.properties import PageSetupProperties
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas

ROOT = Path(__file__).resolve().parent.parent
PAGE_W, PAGE_H = letter
MARGIN = 36

HAZARDOUS_ENERGY = [
    ("he_electrical", "Electrical"), ("he_mechanical", "Mechanical"),
    ("he_thermal", "Thermal"), ("he_pressure", "Pressure"),
    ("he_motion", "Motion"), ("he_gravity", "Gravity"),
    ("he_chemical", "Chemical"), ("he_radiation", "Radiation"),
    ("he_biological", "Biological"), ("he_noise", "Noise"),
]
TASK_HAZARDS = [
    ("tha_loto", "LOTO"), ("tha_confined_space", "Confined Space"),
    ("tha_elevated_work", "Elevated Work"), ("tha_hot_work", "Hot Work"),
    ("tha_vehicles", "Vehicles"), ("tha_mobile_equip", "Mobile Equipment"),
    ("tha_heavy_equip", "Heavy Equipment"), ("tha_power_tools", "Power Tools"),
    ("tha_slips_trips", "Slips / Trips / Falls"), ("tha_nip_pinch", "Nip / Pinch Points"),
    ("tha_break_sys_integ", "Breaking System Integrity"), ("tha_ergonomics", "Ergonomics"),
    ("tha_extreme_heat", "Extreme Heat"), ("tha_extreme_cold", "Extreme Cold"),
    ("tha_high_wind", "High Wind"), ("tha_rain_snow", "Rain / Snow"),
    ("tha_other", "Other"),
]
PPE = [
    ("ppe_hard_hat", "Hard Hat"), ("ppe_safety_glasses", "Safety Glasses"),
    ("ppe_gloves", "Gloves"), ("ppe_fr_clothing", "FR Clothing"),
    ("ppe_arc_flash_suit", "Arc Flash Suit"), ("ppe_arc_flash_hood", "Arc Flash Hood"),
    ("ppe_insul_gloves", "Insulating Gloves"), ("ppe_face_shield", "Face Shield"),
    ("ppe_hearing_prot", "Hearing Protection"), ("ppe_respirator", "Respirator"),
    ("ppe_safety_shoes", "Safety Shoes"), ("ppe_fall_protect", "Fall Protection"),
    ("ppe_cut_protect", "Cut Protection"), ("ppe_welding_hood", "Welding Hood"),
    ("ppe_brine_gear", "Brine Gear"), ("ppe_chem_res_cloth", "Chem-Resistant Clothing"),
    ("ppe_chem_res_gloves", "Chem-Resistant Gloves"), ("ppe_chem_res_boot", "Chem-Resistant Boots"),
    ("ppe_other1", "Other"), ("ppe_other2", "Other"),
]
SAFEGUARDS = [
    ("ts_review_procedures", "Review Procedures"), ("ts_valve_positions", "Verify Valve Positions"),
    ("ts_cross_connects", "Check Cross-Connects"), ("ts_barricade", "Barricade Area"),
    ("ts_adjacent_work", "Adjacent Work Review"), ("ts_haz_comm", "Hazard Communication / SDS"),
    ("ts_addl_permits", "Additional Permits"), ("ts_equip_inspect", "Equipment Inspection"),
    ("ts_high_temp_lines", "High-Temp Line Check"), ("ts_min_approach", "Minimum Approach Distance"),
    ("ts_safety_observer", "Safety Observer"), ("ts_slings_lifting", "Slings / Lifting Devices"),
    ("ts_other1", "Other"),
]
PERMITS = [
    ("permit_loto", "LOTO"), ("permit_hot_work", "Hot Work"),
    ("permit_confined_sp", "Confined Space"), ("permit_excavation", "Excavation / Trenching"),
    ("permit_other", "Other"),
]
PRE_JOB = [
    ("first_time", "First time performing this work?"),
    ("training", "Training complete?"),
    ("inspection", "Pre-job inspection done?"),
]


class Form:
    def __init__(self, path: Path):
        self.c = canvas.Canvas(str(path), pagesize=letter)
        self.c.setLineWidth(0.5)
        self.c.setTitle("Safe Work Permit")
        self.y = PAGE_H - MARGIN

    def heading(self, text):
        self.y -= 18
        self.c.setFont("Helvetica-Bold", 10)
        self.c.setFillGray(0.9)
        self.c.rect(MARGIN, self.y - 3, PAGE_W - 2 * MARGIN, 14, stroke=0, fill=1)
        self.c.setFillGray(0)
        self.c.drawString(MARGIN + 4, self.y, text)
        self.y -= 6

    def text(self, name, label, x, w, h=16, multiline=False):
        self.c.setFont("Helvetica", 7)
        self.c.drawString(x, self.y - 8, label)
        self.c.rect(x, self.y - 10 - h, w, h)
        self.c.acroForm.textfield(
            name=name, x=x, y=self.y - 10 - h, width=w, height=h,
            borderWidth=0.5, fontSize=9,
            fieldFlags="multiline" if multiline else "", maxlen=100000,
        )

    def check(self, name, label, x, y):
        self.c.rect(x, y, 9, 9)
        self.c.acroForm.checkbox(name=name, x=x, y=y, size=9, borderWidth=0.5, buttonStyle="cross")
        self.c.setFont("Helvetica", 8)
        self.c.drawString(x + 12, y + 1.5, label)

    def grid(self, items, cols):
        col_w = (PAGE_W - 2 * MARGIN) / cols
        rows = (len(items) + cols - 1) // cols
        for i, (name, label) in enumerate(items):
            r, col = divmod(i, cols)
            self.check(name, label, MARGIN + col * col_w + 2, self.y - 14 - r * 13)
        self.y -= 14 + rows * 13

    def row(self, fields, h=16):
        """fields: (name, label, width_fraction[, multiline])"""
        x = MARGIN
        total = PAGE_W - 2 * MARGIN
        for f in fields:
            name, label, frac = f[:3]
            w = total * frac - 4
            self.text(name, label, x, w, h, multiline=len(f) > 3 and f[3])
            x += total * frac
        self.y -= 14 + h


def make_swp(path: Path):
    f = Form(path)
    c = f.c
    c.setFont("Helvetica-Bold", 16)
    c.drawString(MARGIN, f.y - 12, "SAFE WORK PERMIT")
    c.setFont("Helvetica", 8)
    c.drawRightString(PAGE_W - MARGIN, f.y - 12, "Form SWP-01")
    f.y -= 22

    f.heading("PERMIT INFORMATION")
    f.row([("company", "Company", 0.35), ("permit_number", "SWP Number", 0.25),
           ("date", "Date", 0.2), ("time", "Time", 0.2)])
    f.row([("location", "Location", 1.0)])
    f.row([("scope_of_work", "Scope of Work", 1.0, True)], h=40)

    f.heading("PRE-JOB QUESTIONS")
    for name, label in PRE_JOB:
        f.y -= 13
        c.setFont("Helvetica", 8)
        c.drawString(MARGIN + 2, f.y + 1.5, label)
        f.check(f"{name}_yes", "Yes", MARGIN + 260, f.y)
        f.check(f"{name}_no", "No", MARGIN + 310, f.y)
    f.y -= 4

    f.heading("HAZARDOUS ENERGY")
    f.grid(HAZARDOUS_ENERGY, 5)

    f.heading("TASK HAZARD ASSESSMENT")
    f.grid(TASK_HAZARDS, 4)
    f.row([("tha_other_txt", "Other hazards", 1.0)])

    f.heading("PERSONAL PROTECTIVE EQUIPMENT")
    f.grid(PPE, 4)

    f.heading("HAZARD MITIGATION (hierarchy of controls)")
    f.row([("hazard_mitigation", "", 1.0, True)], h=40)

    c.showPage()
    c.setLineWidth(0.5)
    f.y = PAGE_H - MARGIN

    f.heading("TASK SAFEGUARDS (pre-job safety meeting)")
    f.grid(SAFEGUARDS, 3)
    f.row([("ts_other1_txt", "Other safeguard", 1.0)])

    f.heading("ASSOCIATED PERMITS")
    for name, label in PERMITS:
        f.y -= 20
        f.check(f"{name}_chk", label, MARGIN + 2, f.y + 3)
        c.setFont("Helvetica", 7)
        c.drawString(MARGIN + 200, f.y + 4, "Permit #")
        c.rect(MARGIN + 235, f.y, 200, 14)
        c.acroForm.textfield(name=name, x=MARGIN + 235, y=f.y, width=200, height=14,
                             borderWidth=0.5, fontSize=9)
    f.y -= 6

    f.heading("ADDITIONAL INSTRUCTIONS OR COMMENTS")
    f.row([("additional_comments", "", 1.0, True)], h=48)

    f.heading("CONTINUATION")
    f.y -= 13
    c.setFont("Helvetica", 8)
    c.drawString(MARGIN + 2, f.y + 1.5, "Continuation sheet required?")
    f.check("continuation_yes", "Yes", MARGIN + 260, f.y)
    f.check("continuation_no", "No", MARGIN + 310, f.y)
    f.y -= 6

    f.heading("PERMIT ISSUED")
    f.row([("issued_date", "Date", 0.2), ("issued_time", "Time", 0.15),
           ("issued_wl_name", "Work Lead", 0.325), ("issued_auth_name", "Authorizer", 0.325)])
    f.heading("PERMIT CLOSED")
    f.row([("closed_date", "Date", 0.2), ("closed_time", "Time", 0.15),
           ("closed_wl_name", "Work Lead", 0.325), ("closed_auth_name", "Authorizer", 0.325)])

    f.heading("ATTENDEES (sign in person)")
    c.setFont("Helvetica", 7)
    f.y -= 10
    c.drawString(MARGIN, f.y, "Name")
    c.drawString(MARGIN + 260, f.y, "Company")
    c.drawString(MARGIN + 440, f.y, "Signature")
    for _ in range(10):
        f.y -= 18
        c.line(MARGIN, f.y, MARGIN + 250, f.y)
        c.line(MARGIN + 260, f.y, MARGIN + 430, f.y)
        c.line(MARGIN + 440, f.y, PAGE_W - MARGIN, f.y)
    c.save()


def make_loto(path: Path):
    thin = Side(style="thin")
    box = Border(left=thin, right=thin, top=thin, bottom=thin)
    bold = Font(name="Times New Roman", size=11, bold=True)
    title = Font(name="Times New Roman", size=16, bold=True)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)

    wb = openpyxl.Workbook()
    ws1 = wb.active
    ws1.title = "Lockout-Tagout Form Part I"
    for col, width in {"A": 2, "B": 16.11, "C": 35.67, "D": 35.67, "E": 34.66, "F": 2, "G": 2,
                       "H": 16.89, "I": 37, "J": 37, "K": 34.66, "L": 34.66}.items():
        ws1.column_dimensions[col].width = width

    ws1.merge_cells("B2:L2")
    ws1["B2"] = "LOCKOUT / TAGOUT — PART I: ISOLATION BOUNDARIES"
    ws1["B2"].font, ws1["B2"].alignment = title, center

    for label_cell, value_range, label in [("B4", "D4:E4", "LOTO #"), ("I4", "K4:L4", "Date"),
                                           ("B5", "D5:L5", "Equipment"), ("B6", "D6:L6", "Job description")]:
        ws1.merge_cells(f"{label_cell}:{chr(ord(label_cell[0]) + 1)}{label_cell[1:]}")
        ws1[label_cell] = label
        ws1[label_cell].font = bold
        ws1.merge_cells(value_range)
    for r in range(4, 9):
        ws1.row_dimensions[r].height = 24.8

    for col, text in [("B", "Device"), ("C", "Equipment / Location"), ("E", "Position"),
                      ("H", "Device"), ("I", "Equipment / Location"), ("K", "Position")]:
        ws1[f"{col}10"] = text
        ws1[f"{col}10"].font, ws1[f"{col}10"].alignment = bold, center
    ws1.merge_cells("C10:D10")
    ws1.merge_cells("I10:J10")
    for r in range(11, 33):
        ws1.row_dimensions[r].height = 39.8
        ws1.merge_cells(f"C{r}:D{r}")
        ws1.merge_cells(f"I{r}:J{r}")
        for col in "BCDEHIJK":
            ws1[f"{col}{r}"].border = box

    ws1["B35"] = "Group leader"
    ws1["B35"].font = bold
    ws1.merge_cells("B36:C36")
    ws1.row_dimensions[36].height = 45
    ws1["B36"].border = box

    ws2 = wb.create_sheet("Lockout-Tagout Form Part II")
    for col, width in {"A": 55.67, "B": 55.33, "C": 15.88, "D": 2, "E": 55.67}.items():
        ws2.column_dimensions[col].width = width
    ws2.merge_cells("A2:E2")
    ws2["A2"] = "LOCKOUT / TAGOUT — PART II: VERIFICATION"
    ws2["A2"].font, ws2["A2"].alignment = title, center
    ws2["A5"], ws2["E5"] = "Person placing lockout", "Date"
    ws2["A5"].font = ws2["E5"].font = bold
    ws2.row_dimensions[6].height = 45
    ws2["A6"].border = ws2["E6"].border = box
    ws2["B17"], ws2["C17"] = "Verified by", "Date"
    ws2["B17"].font = ws2["C17"].font = bold
    for r in range(18, 38):
        ws2.row_dimensions[r].height = 24.8
        ws2[f"B{r}"].border = ws2[f"C{r}"].border = box

    for ws, orientation in ((ws1, "landscape"), (ws2, "portrait")):
        ws.page_setup.orientation = orientation
        ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 1

    wb.save(path)


if __name__ == "__main__":
    (ROOT / "forms" / "swp").mkdir(parents=True, exist_ok=True)
    (ROOT / "forms" / "loto").mkdir(parents=True, exist_ok=True)
    make_swp(ROOT / "forms" / "swp" / "swp-form.pdf")
    make_loto(ROOT / "forms" / "loto" / "loto-form.xlsx")
    print("wrote forms/swp/swp-form.pdf and forms/loto/loto-form.xlsx")
